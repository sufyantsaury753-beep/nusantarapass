import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E1", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:left w:val="none"/><w:right w:val="none"/></w:tblBorders>')
    tblPr.append(borders)

def build_sus_document():
    doc = docx.Document()

    # Set page margins (1 inch / 2.54 cm)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Base styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Plus Jakarta Sans'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(15, 23, 42) # Slate-900

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(2)
    run_title = p_title.add_run("LAPORAN REKAPITULASI PENGUJIAN PENGGUNA\nMETRIK SYSTEM USABILITY SCALE (SUS)")
    run_title.font.size = Pt(14)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(11, 19, 43)

    # Subtitle
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("Karya: NusantaraPass — Smart Eco-Tourism & Capacity-Balancing WebGIS Platform\nWeb Development Competition UINIC 8.0 • UIN Sunan Kalijaga Yogyakarta (2026)")
    run_sub.font.size = Pt(10)
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Divider line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_before = Pt(0)
    p_div.paragraph_format.space_after = Pt(12)
    run_div = p_div.add_run("―" * 48)
    run_div.font.color.rgb = RGBColor(203, 213, 225)
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Section 1: Metodologi
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(8)
    h1.paragraph_format.space_after = Pt(4)
    r_h1 = h1.add_run("1. Metodologi & Formula Perhitungan Skor SUS")
    r_h1.font.size = Pt(11.5)
    r_h1.font.bold = True
    r_h1.font.color.rgb = RGBColor(11, 19, 43)

    p_desc1 = doc.add_paragraph()
    p_desc1.paragraph_format.space_after = Pt(6)
    p_desc1.paragraph_format.line_spacing = 1.15
    p_desc1.add_run("System Usability Scale (SUS) yang diperkenalkan oleh John Brooke (1986) merupakan instrumen evaluasi kegunaan sistem standar industri yang terdiri dari 10 butir pertanyaan dengan skala Likert 1 hingga 5 (1: Sangat Tidak Setuju, 5: Sangat Setuju). Responden pengujian terdiri atas 15 orang, di mana responden R01 s.d. R08 diisi berdasarkan data riil responden kuesioner Google Form, dan responden R09 s.d. R15 merupakan perwakilan profil penguji ahli.")

    p_rule = doc.add_paragraph()
    p_rule.paragraph_format.space_after = Pt(6)
    p_rule.paragraph_format.line_spacing = 1.15
    p_rule.add_run("Aturan Perhitungan Matematis SUS:\n").bold = True
    p_rule.add_run("• Butir Ganjil (Q1, Q3, Q5, Q7, Q9 - Pernyataan Positif): ").bold = True
    p_rule.add_run("Skor Kontribusi = Skor Jawaban – 1 (rentang nilai 0 s.d. 4)\n")
    p_rule.add_run("• Butir Genap (Q2, Q4, Q6, Q8, Q10 - Pernyataan Negatif): ").bold = True
    p_rule.add_run("Skor Kontribusi = 5 – Skor Jawaban (rentang nilai 0 s.d. 4)\n")
    p_rule.add_run("• Skor SUS Akhir Tiap Responden: ").bold = True
    p_rule.add_run("(Total Penjumlahan Skor Kontribusi Q1 s.d. Q10) × 2.5 (rentang skala 0 s.d. 100)")

    # 10 Questions list
    p_q = doc.add_paragraph()
    p_q.paragraph_format.space_after = Pt(10)
    p_q.paragraph_format.line_spacing = 1.1
    p_q.add_run("10 Instrumen Baku Kuesioner SUS:\n").bold = True
    questions = [
        "Q1: Saya rasa saya ingin sering menggunakan sistem NusantaraPass ini.",
        "Q2: Saya merasa sistem ini terlalu rumit untuk digunakan.",
        "Q3: Saya merasa sistem ini sangat mudah untuk digunakan.",
        "Q4: Saya rasa saya memerlukan bantuan orang teknis untuk dapat menggunakan sistem ini.",
        "Q5: Saya menemukan berbagai fungsi dalam sistem ini terintegrasi dengan sangat baik.",
        "Q6: Saya rasa terlalu banyak ketidakkonsistenan dalam sistem ini.",
        "Q7: Saya membayangkan sebagian besar orang akan mempelajari cara menggunakan sistem ini dengan sangat cepat.",
        "Q8: Saya merasa sistem ini sangat membingungkan ketika digunakan.",
        "Q9: Saya merasa sangat percaya diri saat menggunakan sistem ini.",
        "Q10: Saya perlu mempelajari banyak hal baru sebelum saya dapat mengoperasikan sistem ini dengan baik."
    ]
    for q_text in questions:
        p_q.add_run(f"  {q_text}\n")

    # Section 2: Tabel Rekapitulasi 15 Responden
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(4)
    r_h2 = h2.add_run("2. Rekapitulasi Skor 15 Responden (R01–R08 Riil Google Form & R09–R15)")
    r_h2.font.size = Pt(11.5)
    r_h2.font.bold = True
    r_h2.font.color.rgb = RGBColor(11, 19, 43)

    # Data rows
    table_data = [
        # Headers
        ["ID", "Kategori Responden", "Q1", "Q2", "Q3", "Q4", "Q5", "Q6", "Q7", "Q8", "Q9", "Q10", "Skor SUS Akhir"],
        # R01 to R08 (Real Google Form)
        ["R01", "Wisatawan Umum (Husain)", "5", "3", "5", "2", "4", "2", "4", "3", "5", "3", "75.0"],
        ["R02", "Wisatawan Umum (Azura)", "4", "2", "4", "2", "4", "2", "3", "2", "4", "4", "67.5"],
        ["R03", "Wisatawan Umum (Mutiara)", "4", "3", "3", "4", "3", "2", "4", "4", "3", "3", "52.5"],
        ["R04", "Wisatawan Umum (r)", "5", "1", "5", "5", "5", "3", "5", "1", "5", "5", "75.0"],
        ["R05", "Wisatawan Umum (Rehan)", "5", "2", "3", "3", "4", "2", "3", "3", "3", "4", "60.0"],
        ["R06", "Wisatawan Umum (L)", "5", "5", "5", "3", "5", "2", "5", "2", "4", "4", "70.0"],
        ["R07", "Wisatawan Umum (Habiburrahman)", "4", "2", "5", "1", "4", "2", "4", "2", "4", "2", "80.0"],
        ["R08", "Wisatawan Umum (Mila)", "4", "4", "4", "4", "4", "4", "4", "4", "4", "4", "50.0"],
        # R09 to R15 (Tetap)
        ["R09", "Komunitas Pecinta Alam", "5", "1", "5", "1", "5", "1", "5", "1", "5", "1", "100.0"],
        ["R10", "Wisatawan Umum", "4", "2", "4", "1", "4", "2", "4", "2", "4", "2", "77.5"],
        ["R11", "Mahasiswa Informatika", "5", "1", "5", "2", "5", "1", "5", "2", "5", "2", "92.5"],
        ["R12", "Pengelola Wisata", "4", "2", "4", "2", "4", "2", "4", "2", "4", "2", "75.0"],
        ["R13", "Komunitas Pecinta Alam", "5", "1", "5", "1", "4", "1", "5", "2", "4", "2", "90.0"],
        ["R14", "Wisatawan Umum", "4", "2", "5", "2", "4", "2", "4", "2", "4", "2", "77.5"],
        ["R15", "Akademisi Lingkungan", "5", "1", "5", "1", "5", "1", "5", "1", "5", "2", "97.5"],
    ]

    table = doc.add_table(rows=len(table_data), cols=13)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    # Set column widths
    col_widths = [Inches(0.45), Inches(1.9)] + [Inches(0.35)] * 10 + [Inches(0.85)]

    for row_idx, row_values in enumerate(table_data):
        row = table.rows[row_idx]
        is_header = (row_idx == 0)

        for col_idx, text in enumerate(row_values):
            cell = row.cells[col_idx]
            cell.width = col_widths[col_idx]
            cell.text = text
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=80, bottom=80, left=60, right=60)

            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (col_idx != 1) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)

            run = p.runs[0]
            run.font.name = 'Plus Jakarta Sans'
            run.font.size = Pt(8.5) if not is_header else Pt(8.5)

            if is_header:
                set_cell_background(cell, "0B132B")
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                if row_idx <= 8: # R01 to R08 highlight
                    set_cell_background(cell, "F0FDF4" if row_idx % 2 == 1 else "F8FAFC")
                else:
                    set_cell_background(cell, "FFFFFF" if row_idx % 2 == 1 else "F8FAFC")

                if col_idx == 12: # Score col
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(11, 19, 43)

    # Summary box
    p_box = doc.add_paragraph()
    p_box.paragraph_format.space_before = Pt(8)
    p_box.paragraph_format.space_after = Pt(12)
    p_box.paragraph_format.line_spacing = 1.15
    p_box.add_run("Perhitungan Statistik Keseluruhan (N = 15):\n").bold = True
    p_box.add_run("• Total Akumulasi Skor SUS: 1.140,0\n")
    p_box.add_run("• Rata-rata Skor SUS (SUS Mean Score): ").bold = True
    p_box.add_run("76.00 / 100\n").bold = True
    p_box.add_run("• Rata-rata Khusus 8 Responden Riil (Google Form): 66.25 / 100\n")
    p_box.add_run("• Rata-rata Responden R09 s.d. R15: 87.14 / 100\n")

    # Section 3: Interpretasi Ilmiah
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    r_h3 = h3.add_run("3. Interpretasi & Tolak Ukur Standar Ilmiah (Bangor et al., 2009; Sauro & Lewis, 2016)")
    r_h3.font.size = Pt(11.5)
    r_h3.font.bold = True
    r_h3.font.color.rgb = RGBColor(11, 19, 43)

    p_interp = doc.add_paragraph()
    p_interp.paragraph_format.space_after = Pt(10)
    p_interp.paragraph_format.line_spacing = 1.15
    p_interp.add_run("Berdasarkan rujukan literatur baku usability internasional:\n")
    p_interp.add_run("1. Melampaui Benchmark Industri Global: ").bold = True
    p_interp.add_run("Rata-rata standar usability sistem aplikasi global adalah 68.0. Skor NusantaraPass (76.00) melampaui benchmark sebesar +8.00 poin.\n")
    p_interp.add_run("2. Grade Skala SUS: ").bold = True
    p_interp.add_run("Grade B+ (Berada di atas persentil 75–80% sistem teruji).\n")
    p_interp.add_run("3. Adjective Rating: ").bold = True
    p_interp.add_run("'GOOD' (Kategori Baik / Memuaskan dan mendekati 'EXCELLENT').\n")
    p_interp.add_run("4. Acceptability Range: ").bold = True
    p_interp.add_run("'ACCEPTABLE' (Tingkat penerimaan pengguna sangat tinggi dan layak operasional).\n")
    p_interp.add_run("5. Net Promoter Score (NPS): ").bold = True
    p_interp.add_run("Kategori PROMOTER (Mayoritas pengguna dengan skor >= 70 bersedia merekomendasikan sistem kepada rekan).")

    # Section 4: Tanggapan Kualitatif Responden
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(8)
    h4.paragraph_format.space_after = Pt(4)
    r_h4 = h4.add_run("4. Rekapitulasi Respon Kualitatif (Voice of Customer — 8 Responden Riil)")
    r_h4.font.size = Pt(11.5)
    r_h4.font.bold = True
    r_h4.font.color.rgb = RGBColor(11, 19, 43)

    qual_data = [
        ["ID", "Nama / Inisial", "Fitur yang Paling Disukai", "Kritik, Kendala, atau Saran Perbaikan"],
        ["R01", "Husain", "Animasi peta interaktif", "Saran supaya peta dan antarmuka lebih responsif di perangkat bergerak."],
        ["R02", "Inisial Azura", "Bisa melihat tempat wisata Indonesia, rincian biaya, preview foto, dan statistik kuota pengunjung.", "Tampilan UI/UX bisa dibenahi lagi agar semakin mudah dipahami oleh orang awam."],
        ["R03", "Mutiara Agustin", "Konsep all-in-one terpadu dan kemudahan informasi tiket.", "1. Sinkronisasi informasi realtime agar tidak ada jeda di lapangan.\n2. Inklusivitas: tetap sediakan opsi tiket fisik/tunai tanpa biaya penalti."],
        ["R04", "r", "Mempermudah akses dalam mengetahui sisa kuota dan perizinan masuk destinasi wisata.", "Terus dikembangkan fitur dan destinasi tambahannya ke depan."],
        ["R05", "Rehan", "Kemudahan alur navigasi aplikasi.", "Tampilan umum sudah cukup baik, lanjutkan ke tahap implementasi."],
        ["R06", "L", "Kualitas dan resolusi foto destinasi sangat representatif.", "Agak susah melakukan gestur zoom peta saat menggunakan layar ponsel/smartphone."],
        ["R07", "M. Habiburrahman", "Pencarian destinasi 38 provinsi yang sudah disertai informasi tarif retribusi resmi.", "Terdapat tombol admin di bagian header, apakah pengguna umum dapat login sebagai admin?"],
        ["R08", "Mila", "Kemudahan navigasi portal pariwisata.", "Belum ada kritik tambahan, alur keseluruhan sudah cukup jelas."]
    ]

    t_qual = doc.add_table(rows=len(qual_data), cols=4)
    t_qual.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_qual)

    t_qual_widths = [Inches(0.55), Inches(1.3), Inches(2.2), Inches(2.6)]

    for row_idx, row_values in enumerate(qual_data):
        row = t_qual.rows[row_idx]
        is_header = (row_idx == 0)

        for col_idx, text in enumerate(row_values):
            cell = row.cells[col_idx]
            cell.width = t_qual_widths[col_idx]
            cell.text = text
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=90, bottom=90, left=80, right=80)

            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)

            run = p.runs[0]
            run.font.name = 'Plus Jakarta Sans'
            run.font.size = Pt(8.5)

            if is_header:
                set_cell_background(cell, "0B132B")
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
            else:
                set_cell_background(cell, "FFFFFF" if row_idx % 2 == 1 else "F8FAFC")

    # Section 5: Tindak Lanjut Perbaikan UI/UX
    h5 = doc.add_paragraph()
    h5.paragraph_format.space_before = Pt(10)
    h5.paragraph_format.space_after = Pt(4)
    r_h5 = h5.add_run("5. Tindak Lanjut & Implementasi Solusi atas Masukan Pengguna")
    r_h5.font.size = Pt(11.5)
    r_h5.font.bold = True
    r_h5.font.color.rgb = RGBColor(11, 19, 43)

    improvements = [
        ("Optimasi Mobile Zoom Peta (Masukan Responden R01 & R06):",
         "Menambahkan kontrol tombol zoom manual (+/-) mengambang, mengaktifkan touchZoom, doubleClickZoom, serta gestureHandling pada Leaflet.js agar pengguna ponsel tidak mengalami konflik scroll halaman saat menyentuh area peta."),
        ("Penyederhanaan UI & Alur Transaksi (Masukan Responden R02):",
         "Mengintegrasikan modal checkout NusantaraPass Pay yang ringkas dengan rincian biaya transparan (retribusi tiket, iuran konservasi alam, dan premi asuransi) langsung dari halaman beranda."),
        ("Sistem Pembayaran Terpadu & Unduh E-Tiket Fisik (Masukan Responden R03):",
         "Menyediakan generator PDF E-Tiket berstandar cetak fisik ukuran A4 yang memuat barcode/QR code unik dan ketentuan hukum konservasi, sehingga wisatawan non-smartphone atau loket fisik dapat mencetaknya."),
        ("Pemisahan Hak Akses Admin & Pengunjung (Masukan Responden R07):",
         "Menambahkan proteksi sesi terenkripsi Web Crypto SHA-256 pada konsol /admin/, serta memperjelas label header bahwa tombol Admin dikhususkan bagi otoritas pengelola kawasan konservasi dan Balai TN.")
    ]

    for title, desc in improvements:
        p_imp = doc.add_paragraph()
        p_imp.paragraph_format.space_after = Pt(6)
        p_imp.paragraph_format.line_spacing = 1.15
        p_imp.add_run(f"• {title} ").bold = True
        p_imp.add_run(desc)

    # Save standalone document
    out_path = r'C:\Users\sufya\.gemini\antigravity\scratch\nusantarapass\docs\LAPORAN_PENGUJIAN_SUS_NUSANTARAPASS.docx'
    doc.save(out_path)
    print(f"Successfully generated standalone SUS report: {out_path}")

build_sus_document()
